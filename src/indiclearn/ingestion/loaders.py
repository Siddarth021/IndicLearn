"""Concrete loaders: TXT/Markdown, PDF (pypdf), PPTX (python-pptx).

Purpose: turn a file into ``ExtractedPage`` objects (1-based page/slide numbers).
Limitations (all NOT YET EVALUATED on real Telugu course material):
  * PDF: text-layer only. Scanned PDFs yield no text (no OCR). Legacy-encoded Telugu fonts may
    extract as garbage; this needs inspection on real files.
  * PPTX: shape text frames and tables; speaker notes included; images/charts ignored.
  * Markdown: headings become ``section``; markup otherwise left as-is.
"""

from __future__ import annotations

from pathlib import Path
from typing import ClassVar

from indiclearn.domain.documents import ExtractedPage
from indiclearn.ingestion.base import ExtractionError


class TextLoader:
    extensions: ClassVar[tuple[str, ...]] = (".txt", ".md")

    def load(self, path: Path) -> list[ExtractedPage]:
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError as exc:
            raise ExtractionError(f"not valid UTF-8: {exc.reason}") from exc
        pages: list[ExtractedPage] = []
        section: str | None = None
        buf: list[str] = []
        for line in text.splitlines():
            if path.suffix.lower() == ".md" and line.startswith("#"):
                if buf:
                    pages.append(ExtractedPage(page_number=len(pages) + 1, text="\n".join(buf), section=section))
                    buf = []
                section = line.lstrip("#").strip() or None
            buf.append(line)
        if buf:
            pages.append(ExtractedPage(page_number=len(pages) + 1, text="\n".join(buf), section=section))
        return pages


class PdfLoader:
    extensions: ClassVar[tuple[str, ...]] = (".pdf",)

    def load(self, path: Path) -> list[ExtractedPage]:
        try:
            from pypdf import PdfReader
            from pypdf.errors import PyPdfError

            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise ExtractionError("PDF is encrypted")
            return [
                ExtractedPage(page_number=i, text=page.extract_text() or "")
                for i, page in enumerate(reader.pages, start=1)
            ]
        except ExtractionError:
            raise
        except (PyPdfError, ValueError, OSError, KeyError, TypeError) as exc:
            raise ExtractionError(f"PDF could not be parsed: {type(exc).__name__}") from exc


class PptxLoader:
    extensions: ClassVar[tuple[str, ...]] = (".pptx",)

    def load(self, path: Path) -> list[ExtractedPage]:
        try:
            from pptx import Presentation

            prs = Presentation(str(path))
            pages: list[ExtractedPage] = []
            for idx, slide in enumerate(prs.slides, start=1):
                parts: list[str] = []
                title = slide.shapes.title.text.strip() if slide.shapes.title is not None else None
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
                    elif getattr(shape, "has_table", False) and shape.has_table:
                        for row in shape.table.rows:
                            parts.append(" | ".join(cell.text for cell in row.cells))
                if slide.has_notes_slide:
                    parts.append(slide.notes_slide.notes_text_frame.text)
                pages.append(ExtractedPage(page_number=idx, text="\n".join(p for p in parts if p), section=title or None))
            return pages
        except Exception as exc:  # python-pptx raises varied errors on corrupt zips
            raise ExtractionError(f"PPTX could not be parsed: {type(exc).__name__}") from exc
