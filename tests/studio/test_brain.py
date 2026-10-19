"""Hybrid LLM tutor: failover, Telugu/Romanized questions on English notes, reasoning, translation.

No network: a scripted provider stands in for Gemini / Sarvam / Ollama.
"""

from __future__ import annotations

import json
import time

import pytest

from indiclearn.studio import brain as B


class Scripted(B.Provider):
    """Answers like a well-behaved LLM, keyed on which system prompt it receives."""

    def __init__(self, name: str = "fake", fail: bool = False, lang_for: dict[str, str] | None = None) -> None:
        super().__init__(name, "fake-1", 5.0)
        self.fail = fail
        self.calls: list[tuple[str, str]] = []
        self.lang_for = lang_for or {}

    def complete(self, system: str, user: str, json_mode: bool = True, max_tokens: int = 2048) -> str:
        self.calls.append((system[:40], user))
        if self.fail:
            raise RuntimeError("network down")
        if system.startswith("You read a student's question"):
            q = user.split("STUDENT QUESTION:\n", 1)[1]
            lang = next((v for k, v in self.lang_for.items() if k in q), "en")
            return json.dumps({"lang": lang, "romanized": lang == "te" and q.isascii(), "search_query": "What is convolution? convolution integral LTI", "intent": "explain"})
        if system.startswith('You are "Guru"'):
            return "```json\n" + json.dumps({
                "answer": "**కన్వల్యూషన్ (convolution)** అంటే రెండు సిగ్నల్స్‌ను కలిపి కొత్త సిగ్నల్ ఇచ్చే ఆపరేషన్ [p.1]. LTI సిస్టమ్ అవుట్‌పుట్ ఇలాగే వస్తుంది [p.99].",
                "reasoning": ["ప్రశ్న convolution గురించి", "నోట్స్ పేజీ 1 లో నిర్వచనం ఉంది", "LTI తో కలిపాను"],
                "key_terms": [{"term": "convolution", "meaning": "రెండు సిగ్నల్స్ కలయిక"}],
                "follow_ups": ["ఒక ఉదాహరణ ఇవ్వు", "impulse response అంటే?", "క్విజ్ పెట్టు"],
                "check": {"q": "y(t)=x(t)*h(t) లో h(t) ఏమిటి?", "options": ["impulse response", "input", "noise", "gain"], "answer": 0, "why": "h(t) సిస్టమ్ ఇంపల్స్ రెస్పాన్స్"},
                "used_general_knowledge": False, "beyond_notes": ""}, ensure_ascii=False) + "\n```"
        if system.startswith("You translate"):
            d = json.loads(user)
            def tx(v):
                return f"[hi] {v}" if isinstance(v, str) else [tx(x) for x in v] if isinstance(v, list) else v
            return json.dumps({k: tx(v) for k, v in d.items()}, ensure_ascii=False)
        return "{}"


# --------------------------------------------------------------------------- unit
def test_parse_json_tolerates_fences_and_chatter():
    assert B.parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert B.parse_json('Sure! Here it is: {"a": {"b": "x}y"}} hope that helps') == {"a": {"b": "x}y"}}
    assert B.parse_json("no json here") is None
    assert B.parse_json("[1,2]") is None


def test_failover_skips_a_dead_provider_and_cools_it_down():
    dead, live = Scripted("dead", fail=True), Scripted("live")
    br = B.Brain([dead, live])
    out = br.understand("convolution ante enti", [], "en")
    assert out and out["engine"] == "live"
    assert dead.down_until > time.time()  # skipped on the next call ...
    br.understand("again", [], "en")
    assert len(dead.calls) == 1  # ... without another timeout


def test_brain_off_and_no_keys_means_no_providers(monkeypatch):
    for k in ("INDICLEARN_GEMINI_API_KEY", "INDICLEARN_GEMINI_API_KEY_2", "GEMINI_API_KEY", "GOOGLE_API_KEY", "INDICLEARN_SARVAM_API_KEY", "SARVAM_API_KEY",
              "INDICLEARN_OPENAI_BASE_URL", "INDICLEARN_ANTHROPIC_API_KEY", "ANTHROPIC_API_KEY", "INDICLEARN_LLM_BACKEND"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("INDICLEARN_OLLAMA", "on")
    monkeypatch.setenv("INDICLEARN_OLLAMA_URL", "http://127.0.0.1:9")  # nothing listens here
    assert not B.Brain.from_env(ollama_url="http://127.0.0.1:9").available  # no key, no Ollama -> offline extractive
    monkeypatch.setenv("INDICLEARN_GEMINI_API_KEY", "x")
    monkeypatch.setenv("INDICLEARN_BRAIN", "off")
    assert not B.Brain.from_env().available
    monkeypatch.setenv("INDICLEARN_BRAIN", "auto")
    b = B.Brain.from_env(ollama_url="http://127.0.0.1:9")
    # hosted first, local Ollama always kept behind it as the offline fallback
    # one Gemini provider per free-tier model (each has its own daily quota), then Ollama
    names = [p.name for p in b.providers]
    assert names[-1] == "ollama" and set(names[:-1]) == {"gemini"} and b.providers[0].model == "gemini-3.8-flash"
    monkeypatch.setenv("INDICLEARN_GEMINI_MODEL", "gemini-3.8-flash")
    assert [p.name for p in B.Brain.from_env(ollama_url="http://127.0.0.1:9").providers] == ["gemini", "ollama"]
    assert b.available


def test_ollama_picks_the_best_installed_local_model(monkeypatch):
    p = B.OllamaProvider("ollama", "gemma3:4b", 5.0)
    monkeypatch.setattr(p, "_probe", lambda: None)
    p._installed = ("llama3.2:3b", "gemma4:e2b")
    assert p.pick_model() == "gemma4:e2b"  # configured gemma3 is not pulled -> best preferred one
    p._installed = ("gemma3:4b", "gemma4:e2b")
    assert p.pick_model() == "gemma3:4b"  # the configured model wins when present
    p._installed = ("mistral:latest",)
    assert p.pick_model() == "mistral:latest"  # anything beats nothing
    p._installed = ()
    assert p.pick_model() is None and not p.ok()


def test_local_models_get_a_shorter_prompt(en_graph, en_index):
    from indiclearn.studio.tutor import answer

    class LocalScripted(Scripted):
        local = True

    p = LocalScripted(lang_for={"కన్వల్యూషన్": "te"})
    p.local = True
    rep = answer("కన్వల్యూషన్ అంటే ఏమిటి?", "en", en_index, en_graph.nodes, en_graph.edges, "S&S", "en", models=_Models(B.Brain([p])))
    assert rep.lang == "te"
    sys_prompt = next(s for s, _u in p.calls if s.startswith('You are "Guru"'))
    assert sys_prompt  # system captured (first 40 chars)
    tutor_user = next(u for s, u in p.calls if s.startswith('You are "Guru"'))
    assert tutor_user.count("[part ") + tutor_user.count("[p.") <= 3 + 1


# --------------------------------------------------------------------------- tutor
class _Models:
    def __init__(self, br):
        self.brain, self.translator, self.llm = br, None, None


def test_telugu_question_on_english_notes_gets_a_detailed_telugu_answer(en_graph, en_index):
    from indiclearn.studio.tutor import answer

    p = Scripted(lang_for={"కన్వల్యూషన్": "te"})
    rep = answer("కన్వల్యూషన్ అంటే ఏమిటి?", "en", en_index, en_graph.nodes, en_graph.edges, "S&S", "en", models=_Models(B.Brain([p])),
                 history=[{"role": "user", "content": "hi"}])
    assert rep.lang == "te"
    assert "కన్వల్యూషన్" in rep.text and rep.grounded
    assert len(rep.extra["reasoning"]) == 3 and rep.extra["follow_ups"] and rep.extra["check"]["answer"] == 0
    assert rep.citations and all(c["page"] for c in rep.citations)
    assert "[p.99]" not in rep.text  # invented page tag removed
    assert rep.evidence  # the map still learns the student asked about it
    # the tutor saw the English rewrite of the question, not only the Telugu text
    tutor_prompt = next(u for s, u in p.calls if s.startswith('You are "Guru"'))
    assert "convolution" in tutor_prompt.lower() and "hi" in tutor_prompt


def test_romanized_telugu_is_answered_in_telugu_script(en_graph, en_index):
    from indiclearn.studio.tutor import answer

    # even if the model mislabels it "en", our detector catches Romanized Telugu
    p = Scripted(lang_for={})
    rep = answer("convolution ante enti, konchem simple ga cheppandi", "en", en_index, en_graph.nodes, en_graph.edges, "S&S", "en",
                 models=_Models(B.Brain([p])))
    assert rep.lang == "te"


def test_dead_llm_falls_back_to_the_extractive_tutor(en_graph, en_index):
    from indiclearn.studio.tutor import answer

    rep = answer("What is convolution?", "en", en_index, en_graph.nodes, en_graph.edges, "S&S", "en", models=_Models(B.Brain([Scripted(fail=True)])))
    assert rep.grounded and rep.citations
    assert "network down" in rep.extra["notice"]  # the student is told WHY the smart answer is missing


# --------------------------------------------------------------------------- api
@pytest.fixture()
def smart_client(client):
    svc = client.app.state.service
    svc.models.brain = B.Brain([Scripted(lang_for={"కన్వల్యూషన్": "te"})])
    return client, svc


def _ready_doc(client, auth):
    d = client.post("/api/docs/sample/en", headers=auth).json()
    for _ in range(100):
        if client.get(f"/api/docs/{d['id']}", headers=auth).json()["status"] == "ready":
            return d["id"]
        time.sleep(0.05)
    raise AssertionError("doc never became ready")


def test_chat_returns_reasoning_and_history_keeps_it(smart_client, auth):
    client, _svc = smart_client
    doc = _ready_doc(client, auth)
    r = client.post(f"/api/docs/{doc}/chat", json={"message": "కన్వల్యూషన్ అంటే ఏమిటి?", "lang": "en"}, headers=auth).json()
    assert r["lang"] == "te" and r["reasoning"] and r["follow_ups"] and r["check"]
    hist = client.get(f"/api/docs/{doc}/chat", headers=auth).json()
    assert hist[-1]["meta"]["reasoning"] == r["reasoning"]


def test_translate_endpoint_and_language_list(smart_client, auth):
    client, _svc = smart_client
    langs = {x["code"] for x in client.get("/api/languages").json()}
    assert {"te", "hi", "ta", "kn", "te-Latn"} <= langs
    r = client.post("/api/translate", json={"target": "hi", "text": "**Convolution** [p.1]", "reasoning": ["step"]}, headers=auth)
    assert r.status_code == 200 and r.json()["text"].startswith("[hi]") and r.json()["reasoning"] == ["[hi] step"]
    assert client.post("/api/translate", json={"target": "xx", "text": "a"}, headers=auth).status_code == 422


def test_translate_and_tts_explain_themselves_without_models(client, auth):
    r = client.post("/api/translate", json={"target": "ta", "text": "hello"}, headers=auth)
    assert r.status_code == 503 and "Could not translate" in r.json()["detail"] and ".env" in r.json()["detail"]
    r = client.post("/api/tts", json={"text": "నమస్తే", "lang": "te"}, headers=auth)
    assert r.status_code == 503 and "mms_tts_tel" in r.json()["detail"]


# --------------------------------------------------------------------------- knowledge graph
def test_cleanup_drops_modifiers_merges_duplicates_and_keeps_prereqs_forward():
    from pathlib import Path

    from indiclearn.studio.kgraph import build_knowledge_graph
    from indiclearn.studio.kgraph_refine import clean_graph, reading_order

    nodes = [
        {"id": "root", "label": "Doc", "type": "root", "level": 0, "parent": None, "page": 1, "order": 0},
        {"id": "s1", "label": "Methods", "type": "chapter", "level": 1, "parent": "root", "page": 1, "order": 1, "is_term": True},
        {"id": "a", "label": "Self-supervised", "type": "concept", "parent": "s1", "section_id": "s1", "page": 1, "order": 2, "importance": 0.4},
        {"id": "b", "label": "Student model", "type": "concept", "parent": "s1", "section_id": "s1", "page": 1, "order": 3, "importance": 0.4},
        {"id": "c", "label": "Student network", "type": "concept", "parent": "s1", "section_id": "s1", "page": 1, "order": 4, "importance": 0.6},
        {"id": "d", "label": "Teacher", "type": "concept", "parent": "s1", "section_id": "s1", "page": 1, "order": 5, "importance": 0.5},
    ]
    edges = [{"id": "e1", "source": "d", "target": "b", "type": "prereq"},   # teacher before student: forward, kept (b->c merged)
             {"id": "e2", "source": "c", "target": "d", "type": "prereq"},   # backward: dropped
             {"id": "e3", "source": "d", "target": "s1", "type": "prereq"}]  # into a chapter: dropped
    text = "The Teacher is trained first. Self-supervised pretraining and self-supervised learning help. A Student model is small. The Student network copies."
    kept, new_edges, rep = clean_graph(nodes, edges, text)
    ids = {n["id"] for n in kept}
    assert "a" not in ids and rep["dropped_modifiers"] == 1
    assert rep["merged_duplicates"] == 1 and len({"b", "c"} & ids) == 1
    pre = [(e["source"], e["target"]) for e in new_edges if e["type"] == "prereq"]
    assert len(pre) == 1 and pre[0][0] == "d"
    assert reading_order(kept, text)["d"] < reading_order(kept, text)[pre[0][1]]
    # the real sample graph keeps its sensible chain
    g = build_knowledge_graph(__import__("indiclearn.studio.parse", fromlist=["x"]).parse_file(
        Path(__file__).resolve().parents[2] / "data/samples/signals_and_systems_notes.md", "n.md"))
    from indiclearn.studio.kgraph_refine import refine
    refine(g)
    lab = {n["id"]: n["label"] for n in g.nodes}
    assert ("Fourier Series", "Fourier Transform") in {(lab[e["source"]], lab[e["target"]]) for e in g.edges if e["type"] == "prereq"}


class GraphLLM(B.Provider):
    def __init__(self):
        super().__init__("graphfake", "g", 5.0)

    def complete(self, system, user, json_mode=True, max_tokens=2048):
        return json.dumps({"concepts": [
            {"name": n, "aliases": a, "section": "nope", "definition": f"{n} defined.", "importance": 0.7}
            for n, a in [("Convolution", []), ("Impulse response", ["h(t)"]), ("Fourier Transform", ["FT"]), ("Laplace Transform", []),
                         ("Sampling theorem", ["Nyquist"]), ("Aliasing", []), ("Quantum gravity", [])]],
            "prereqs": [{"before": "Impulse response", "after": "Convolution", "why": "y = x*h"},
                        {"before": "Convolution", "after": "Impulse response"},  # would close a cycle
                        {"before": "Fourier Transform", "after": "Laplace Transform"}, {"before": "Ghost", "after": "Aliasing"}],
            "related": [{"a": "Sampling theorem", "b": "Aliasing"}]})


def test_llm_graph_is_validated_against_the_text(en_graph):
    import copy

    from indiclearn.studio.kgraph_refine import refine

    kg = copy.deepcopy(en_graph)
    old_ids = {n["id"] for n in kg.nodes}
    rep = refine(kg, B.Brain([GraphLLM()]))
    assert rep["graph_engine"] == "llm:graphfake"
    labels = {n["label"]: n for n in kg.nodes}
    assert "Quantum gravity" not in labels  # invented: not in the notes
    assert all(n["parent"] for n in kg.nodes if n["type"] == "concept")
    lab = {n["id"]: n["label"] for n in kg.nodes}
    pre = {(lab[e["source"]], lab[e["target"]]) for e in kg.edges if e["type"] == "prereq"}
    assert ("Impulse response", "Convolution") in pre and ("Convolution", "Impulse response") not in pre
    assert any(e.get("why") == "y = x*h" for e in kg.edges)
    # ids of concepts that existed before are reused, so student progress survives
    conv = next(n for n in kg.nodes if n["label"] == "Convolution")
    assert conv["id"] in old_ids
    assert kg.path and kg.path[0] != "root"


def test_streak_and_xp_follow_the_evidence_log(client, auth):
    s0 = client.get("/api/me/stats", headers=auth).json()
    assert s0["xp"] == 0 and s0["streak"] == 0 and s0["level"] == 1
    doc = _ready_doc(client, auth)
    client.post(f"/api/docs/{doc}/chat", json={"message": "What is convolution?", "lang": "en"}, headers=auth)
    s1 = client.get("/api/me/stats", headers=auth).json()
    assert s1["xp"] > 0 and s1["streak"] == 1 and s1["active_today"]


def test_ollama_http_contract_with_a_fake_daemon():
    """Talks HTTP to a tiny fake Ollama: /api/tags for discovery, /api/chat with JSON mode."""
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer

    seen: list[dict] = []

    class H(BaseHTTPRequestHandler):
        def log_message(self, *a):  # keep test output clean
            pass

        def _send(self, obj, code=200):
            b = json.dumps(obj).encode()
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(b)))
            self.end_headers()
            self.wfile.write(b)

        def do_GET(self):
            self._send({"models": [{"name": "gemma4:e2b"}]})

        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            seen.append(body)
            if "think" in body and len(seen) == 1:  # pretend to be an old Ollama that rejects "think"
                return self._send({"error": "unknown field think"}, 400)
            self._send({"message": {"role": "assistant", "content": json.dumps({"lang": "te", "search_query": "convolution", "intent": "explain"})}})

    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        p = B.OllamaProvider("ollama", "gemma3:4b", 5.0, url=f"http://127.0.0.1:{srv.server_port}")
        br = B.Brain([p])
        assert br.available and br.status()["offline_llm"] and "gemma4:e2b" in br.status()["chain"][0]
        u = br.understand("convolution ante enti", [], "en")
        assert u["lang"] == "te" and u["engine"] == "ollama"
        assert seen[-1]["model"] == "gemma4:e2b" and seen[-1]["format"] == "json" and "think" not in seen[-1]
        assert seen[-1]["options"]["num_ctx"] >= 4096
    finally:
        srv.shutdown()


# --------------------------------------------------------------------------- offline models
def test_light_indictrans_processor_round_trips_scripts():
    from indiclearn.studio.neural import LightIndicProcessor, brahmic_convert

    te = "కన్వల్యూషన్ అంటే ఏమిటి? రెండు సిగ్నల్స్।"
    deva = brahmic_convert(te, "Telu", "Deva")
    assert deva.startswith("कन्वल्यूषन्") and brahmic_convert(deva, "Deva", "Telu") == te
    assert "।" in deva  # danda is shared, never shifted into an unassigned code point
    ip = LightIndicProcessor()
    pre = ip.preprocess_batch([te], "tel_Telu", "eng_Latn")[0]
    assert pre.startswith("tel_Telu eng_Latn ") and "क" in pre  # tags + Devanagari, as IndicTrans2 was trained
    assert ip.preprocess_batch(["What is it?"], "eng_Latn", "tel_Telu")[0].startswith("eng_Latn tel_Telu What is it")
    assert ip.postprocess_batch(["कन्वल्यूषन् अंटे एमिटि ?"], "tel_Telu") == ["కన్వల్యూషన్ అంటే ఏమిటి?"]


def test_nllb_covers_every_convert_target_except_romanized():
    from indiclearn.studio.brain import TARGET_LANGS
    from indiclearn.studio.llm import NLLB_CODES

    assert {c for c in TARGET_LANGS if "-" not in c} <= set(NLLB_CODES)


def test_asr_prefers_the_telugu_model_for_telugu(monkeypatch, tmp_path):
    from indiclearn.studio import neural

    for name in ("whisper_te_small", "whisper_small"):
        (tmp_path / "pretrained" / name).mkdir(parents=True)
        (tmp_path / "pretrained" / name / "config.json").write_text("{}")
    monkeypatch.setenv("INDICLEARN_MODELS_DIR", str(tmp_path))
    loaded: list[str] = []

    class FakePipe:
        def __init__(self, model):
            loaded.append(model.split("/")[-1].split("\\\\")[-1])
            self.model = type("M", (), {"config": type("C", (), {})(), "generation_config": type("G", (), {})()})()

    monkeypatch.setattr(neural, "_hf_pipeline", lambda task, model, device=-1: FakePipe(model))
    neural.reset_cache()
    try:
        te = neural.load_recogniser("te")
        en = neural.load_recogniser("en")
        assert te.fixed_lang == "te" and en.fixed_lang == ""
        assert [x.rsplit("/", 1)[-1] for x in loaded] == ["whisper_te_small", "whisper_small"]
        assert neural.status()["asr_te"] == "whisper_te_small"
    finally:
        neural.reset_cache()


# --------------------------------------------------------------------------- round 3: broken text, OCR, summaries, translation
def test_think_tags_and_json_from_reasoning_models_are_handled():
    raw = '<think>let me {think} about it</think>\n```json\n{"answer": "ok"}\n```'
    assert B.parse_json(raw) == {"answer": "ok"}
    assert B.strip_think("<think>never closed {\"a\": 1}") == ""
    assert B.strip_think("reasoning</think>final") == "final"


def test_openai_compat_retries_without_response_format(monkeypatch):
    calls = []

    def fake(url, body, headers, timeout):
        calls.append(dict(body))
        if "response_format" in body:
            raise RuntimeError("HTTP 400: unknown field response_format")
        return {"choices": [{"message": {"content": '{"x": 1}'}}]}

    monkeypatch.setattr(B, "_http_json", fake)
    p = B.OpenAICompatProvider("sarvam", "sarvam-105b", 5, key="k", base="https://x/v1")
    assert B.parse_json(p.complete("s", "u")) == {"x": 1}
    assert "response_format" not in calls[-1] and len(calls) == 2


def test_failed_provider_reports_a_human_reason():
    class Boom(B.Provider):
        def complete(self, *a, **k):
            raise RuntimeError("HTTP 401: invalid api key")

    br = B.Brain([Boom("sarvam", "m")])
    assert br._run("s", "u") is None
    assert "key was rejected" in br.status()["errors"]["sarvam"]
    assert "sarvam" in br.why_unavailable()


def test_pdf_text_damage_is_repaired():
    from indiclearn.studio.parse import _build_vocab, _fix_line

    txt = "can effectively handle data. introduction related work"
    v = _build_vocab([txt])
    assert _fix_line("I NTRODUCTION", v) == "INTRODUCTION"
    assert _fix_line("II. R ELATED W ORK", v) == "II. RELATED WORK"
    assert _fix_line("can e ffectively handle data , , , , .", v) == "can effectively handle data."
    assert _fix_line("a boy in to the room", v) == "a boy in to the room"  # real words are never glued


def test_text_inside_pictures_is_read_with_ocr(tmp_path):
    from indiclearn.studio import ocr

    if ocr.get_engine() is None:
        pytest.skip("no OCR engine installed")
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from pptx.util import Inches

    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 56)
    except OSError:
        pytest.skip("no TrueType font")
    im = Image.new("RGB", (1280, 720), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 40), "Graph Neural Networks", font=font, fill="black")
    d.text((80, 250), "Nodes exchange messages with neighbours in every layer", font=ImageFont.truetype("DejaVuSans.ttf", 34), fill="black")
    png = tmp_path / "s.png"
    im.save(png)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(tmp_path / "pic.pptx")
    im.save(tmp_path / "scan.pdf")
    from indiclearn.studio.parse import parse_file

    for name in ("pic.pptx", "scan.pdf"):
        doc = parse_file(tmp_path / name)
        text = (doc.title + " " + doc.full_text).lower()
        assert "neighbours" in text and "graph neural" in text, (name, text)
        assert any("OCR" in w for w in doc.warnings)


def test_text_inside_grouped_shapes_is_read(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    sl.shapes.title.text = "Slide with a group"
    g = sl.shapes.add_group_shape()
    g.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1)).text_frame.text = "Text hidden inside a group shape"
    prs.save(tmp_path / "g.pptx")
    from indiclearn.studio.parse import parse_file

    doc = parse_file(tmp_path / "g.pptx")
    assert "hidden inside a group" in doc.full_text
    assert sum("Slide with a group" in b.text for b in doc.blocks) == 1  # the title is not duplicated


def test_summary_request_gets_the_whole_topic_and_a_summary_prompt(en_graph, en_index):
    from indiclearn.studio.tutor import answer

    p = Scripted(lang_for={})
    answer("summarize convolution", "en", en_index, en_graph.nodes, en_graph.edges, "S&S", "en", models=_Models(B.Brain([p])))
    assert any("SUMMARY REQUEST" in u for _s, u in p.calls)


def test_extractive_summary_skips_linking_sentences_and_citation_residue():
    from indiclearn.studio.concepts import summarize_sentences

    out = summarize_sentences([
        "Additionally, GNNs can effectively handle data with complex relationships , , , , .",
        "A graph neural network is a model that learns from nodes and the edges that connect them.",
        "Message passing lets every node collect information from its neighbours in each layer.",
        "However, deep stacks of layers can blur node features together.",
    ], 2)
    assert all(not s.startswith(("Additionally", "However")) for s in out) and all(", ," not in s for s in out)


def test_translation_between_two_indian_languages_pivots_through_english(monkeypatch):
    from indiclearn.studio import neural
    from indiclearn.studio.service import StudioService as Service

    class Fake:
        en_indic, indic_en = object(), object()

        def can(self, s, t):
            return neural.IndicTranslator.can(self, s, t)

        def translate(self, texts, s, t, strict=False):
            return [f"[{s}>{t}] {x}" for x in texts]

    loader = lambda: Fake()
    loader.cache_clear = lambda: None  # the test fixture resets caches on teardown
    monkeypatch.setattr(neural, "load_indictrans", loader)
    assert neural.IndicTranslator.can(Fake(), "te", "hi") and neural.IndicTranslator.can(Fake(), "en", "ta")
    svc = Service.__new__(Service)
    svc.models = type("M", (), {"translator": None, "brain": None})()
    out, why = Service._translate_local(svc, {"text": "కన్వల్యూషన్ అంటే ఏమిటి"}, "hi", "te")
    assert out and out["text"].startswith("[te>hi]") and not why
    out, why = Service._translate_local(svc, {"text": "What is convolution"}, "ta", "en")
    assert out and out["text"].startswith("[en>ta]")


def test_english_words_inside_a_telugu_reply_are_read_by_the_english_voice(monkeypatch):
    import io
    import wave

    from indiclearn.studio import neural
    from indiclearn.studio.service import StudioService, join_wavs, split_script_runs

    runs = split_script_runs("కన్వల్యూషన్ అంటే Graph Neural Networks, convolution ante enti?")
    assert [k for k, _ in runs][:2] == ["native", "en"] and runs[1][1].startswith("Graph Neural Networks")

    def tone(rate):
        buf = io.BytesIO()
        with wave.open(buf, "wb") as w:
            w.setnchannels(1), w.setsampwidth(2), w.setframerate(rate)
            w.writeframes(b"\x01\x00" * (rate // 10))
        return buf.getvalue()

    spoken: list[tuple[str, str]] = []

    class V:
        def __init__(self, name, rate):
            self.name, self.rate = name, rate

        def wav(self, text, *a):
            spoken.append((self.name, text))
            return tone(self.rate)

    voices = {"te": V("te", 16000), "en": V("en", 22050)}
    def fake_speaker(lang):
        return voices.get(lang)

    def fake_parler():
        return None

    fake_speaker.cache_clear = fake_parler.cache_clear = lambda: None  # reset_cache() calls these on teardown
    monkeypatch.setattr(neural, "load_speaker", fake_speaker)
    monkeypatch.setattr(neural, "load_parler", fake_parler)
    svc = StudioService.__new__(StudioService)
    out = svc.speak("కన్వల్యూషన్ అంటే Graph Neural Networks", "te")
    assert ("en", "Graph Neural Networks") in spoken and any(n == "te" for n, _ in spoken)
    with wave.open(io.BytesIO(out)) as w:
        assert w.getframerate() == 16000 and w.getnframes() > 3000  # both pieces joined, resampled
    assert join_wavs([tone(16000)])  # single piece is fine


def test_second_gemini_account_takes_over_when_the_first_is_out_of_quota(monkeypatch):
    from indiclearn.studio import brain as B

    for k in ("INDICLEARN_SARVAM_API_KEY", "SARVAM_API_KEY", "GEMINI_API_KEY", "GOOGLE_API_KEY"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("INDICLEARN_BRAIN", "gemini")
    monkeypatch.setenv("INDICLEARN_GEMINI_API_KEY", "key-A")
    monkeypatch.setenv("INDICLEARN_GEMINI_API_KEY_2", "key-B")
    monkeypatch.setenv("INDICLEARN_GEMINI_MODEL", "m1,m2")
    b = B.Brain.from_env()
    assert [(p.model, p.key) for p in b.providers] == [("m1", "key-A"), ("m1", "key-B"), ("m2", "key-A"), ("m2", "key-B")]
    calls = []

    def fake(self, system, user, json_mode=True, max_tokens=2048):
        calls.append((self.model, self.key))
        if self.key == "key-A":
            raise RuntimeError("HTTP 429: quota exceeded GenerateRequestsPerDay")
        return '{"ok": true}'

    monkeypatch.setattr(B.GeminiProvider, "complete", fake)
    assert b._run("s", "u")[0] == {"ok": True}
    assert calls == [("m1", "key-A"), ("m1", "key-B")]  # account A was out of quota: account B answered
    calls.clear()
    b._run("s", "u")
    assert calls == [("m1", "key-B")]  # A is now cooling down - not retried every time
    assert "gemini#2:m1" in b.status()["chain"]


def test_gemini_keys_accept_a_comma_list_and_skip_duplicates(monkeypatch):
    from indiclearn.studio import brain as B

    monkeypatch.setenv("INDICLEARN_GEMINI_API_KEY", "a, b")
    monkeypatch.setenv("INDICLEARN_GEMINI_API_KEY_2", "b")
    monkeypatch.setenv("INDICLEARN_GEMINI_API_KEY_3", "c")
    assert B.gemini_keys() == ["a", "b", "c"]
