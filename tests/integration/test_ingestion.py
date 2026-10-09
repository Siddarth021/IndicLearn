from pathlib import Path

from indiclearn.domain.documents import DocumentStatus


def _pdf(path: Path, pages: list[str]) -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    for t in pages:
        c.drawString(72, 750, t)
        c.showPage()
    c.save()


def _pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Convolution"
    s.placeholders[1].text = "Convolution combines an input with an impulse response."
    s.notes_slide.notes_text_frame.text = "Speaker note about LTI systems."
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Fourier Transform"
    prs.save(str(path))


def test_txt_telugu_code_mixed(service, tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("కన్వల్యూషన్ (convolution) అంటే రెండు సంకేతాల కలయిక.\n\nLTI system గురించి.", encoding="utf-8")
    r = service.ingest(f, "sig")
    assert r.ok and r.chunks
    assert "కన్వల్యూషన్" in r.chunks[0].text and "convolution" in r.chunks[0].text


def test_markdown_sections(service, tmp_path):
    f = tmp_path / "n.md"
    f.write_text("# Signals\nAn energy signal has finite energy.\n# Systems\nCausal systems.", encoding="utf-8")
    r = service.ingest(f, "sig")
    assert r.ok and [p.section for p in r.pages] == ["Signals", "Systems"]


def test_pdf_pages_and_page_numbers(service, tmp_path):
    f = tmp_path / "book.pdf"
    _pdf(f, ["Energy signals page one", "Power signals page two"])
    r = service.ingest(f, "sig")
    assert r.ok and r.document.page_count == 2
    assert {c.page_number for c in r.chunks} == {1, 2}
    assert any("Power signals" in c.text and c.page_number == 2 for c in r.chunks)


def test_pptx_slides_titles_notes(service, tmp_path):
    f = tmp_path / "deck.pptx"
    _pptx(f)
    r = service.ingest(f, "sig")
    assert r.ok and r.pages[0].section == "Convolution" and r.pages[1].section == "Fourier Transform"
    assert "Speaker note" in r.pages[0].text


def test_corrupted_pdf_recorded_not_raised(service, tmp_path):
    f = tmp_path / "bad.pdf"
    f.write_bytes(b"%PDF-1.4 this is not really a pdf \x00\xff")
    r = service.ingest(f, "sig")
    assert r.document.status == DocumentStatus.FAILED and "bad.pdf" == r.document.filename and r.document.error


def test_corrupted_pptx(service, tmp_path):
    f = tmp_path / "bad.pptx"
    f.write_bytes(b"not a zip")
    assert service.ingest(f, "sig").document.status == DocumentStatus.FAILED


def test_empty_document_fails(service, tmp_path):
    f = tmp_path / "empty.txt"
    f.write_text("  \n\n ", encoding="utf-8")
    r = service.ingest(f, "sig")
    assert r.document.status == DocumentStatus.FAILED and "no extractable text" in r.document.error


def test_image_only_pdf_fails_gracefully(service, tmp_path):
    from reportlab.pdfgen import canvas

    f = tmp_path / "blank.pdf"
    c = canvas.Canvas(str(f)); c.showPage(); c.save()
    assert service.ingest(f, "sig").document.status == DocumentStatus.FAILED


def test_duplicate_detection_per_subject(service, tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("Same content here.", encoding="utf-8")
    assert service.ingest(f, "sig").ok
    assert service.ingest(f, "sig").document.status == DocumentStatus.DUPLICATE
    assert service.ingest(f, "other-subject").ok  # same file, different subject is allowed


def test_unsupported_and_missing_file(service, tmp_path):
    f = tmp_path / "x.docx"
    f.write_text("x")
    assert service.ingest(f, "s").document.status == DocumentStatus.FAILED
    assert service.ingest(tmp_path / "ghost.txt", "s").document.status == DocumentStatus.FAILED


def test_non_utf8_text_fails(service, tmp_path):
    f = tmp_path / "l.txt"
    f.write_bytes("é".encode("latin-1"))
    r = service.ingest(f, "s")
    assert r.document.status == DocumentStatus.FAILED and "UTF-8" in r.document.error


def test_size_limit(settings, tmp_path):
    from indiclearn.chunking import RecursiveCharChunker
    from indiclearn.ingestion.service import IngestionService

    settings.ingestion.max_file_mb = 0
    svc = IngestionService(settings, RecursiveCharChunker())
    f = tmp_path / "a.txt"
    f.write_text("hello", encoding="utf-8")
    assert "limit" in svc.ingest(f, "s").document.error
